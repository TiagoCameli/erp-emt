"""Extrai o Lote 09 (CT 00615/2025, DNIT) da planilha oficial e gera o staging e o
esperado da carga (Medição de Contratos, Fase 2).

Uso:
  python3 scripts/migracao-medicao/gerar_carga_lote09.py [caminho_do_xlsx]

Sem argumento, usa a cópia local de trabalho (XLSX_PADRAO). O script recusa qualquer
arquivo cujo sha256 não seja o combinado com o Tiago (planilha oficial, aba "Planilha de
Medição").

Saída (em _retrato/, fora do git: o repositório é público):
  staging_l09.json   {"contrato", "linhas", "medicoes", "quantidades"} para a migration
                      de preparo/carga (Task 2/3) carregar em legado.carga_mc_l09.
  esperado_l09.json  os números da origem, para a conferência da migration de carga.

Decisões do Tiago (26/09/2026), ver docs/superpowers/plans/2026-09-26-medicao-contratos-
fase2-carga-lote09.md, seção "Fonte e decisões":
  - Regra de arredondamento: sem_arredondar (soma exata por Decimal; só arredonda em
    2 casas na saída de grupo e total, no fim da conta).
  - Linha 20 (DOPE) repete no xlsx o código 02.02 da linha 19: entra com 02.02.01. A
    observação do contrato registra a troca.
  - Linhas 184 a 199 (03.16.01 a 03.16.09, ocultas): preço preenchido, quantidade prevista
    vazia. Vazio em serviço vira quantidade "0", como a importação faria.
  - Unidade "un " (linhas 100 e 101): gravada aparada.
  - Grupo 08 (linha 278) tem preço 0 e é título mesmo assim (código de 2 dígitos).
  - Saldo = previsto - acumulado pela conta direta (o módulo não trunca por item como a
    planilha faz na coluna AV).
  - Os 10 períodos (NFs no ERP) e os dados do contrato vêm do vault/PDF do contrato, não
    da planilha: ver CONTRATO e MEDICOES abaixo.
"""
import hashlib
import json
import os
import re
import sys
from decimal import ROUND_HALF_UP, Decimal

import openpyxl

D = os.path.dirname(os.path.abspath(__file__))
RETRATO = os.path.join(D, '_retrato')

# Cópia local de trabalho (fora do git). O script recusa qualquer arquivo com outro hash.
XLSX_PADRAO = '/Users/tiagocameli/.claude/jobs/84e7ccda/tmp/fase2/v12.xlsx'
SHA256_ESPERADO = '2a29e7473cca3e057a700a91c58d8d992e89ea42e2c74a094d1f3e59121fcb0b'

ABA = 'Planilha de Medição'
PRIMEIRA_LINHA = 15
ULTIMA_LINHA = 279  # inclusive; a linha 280 é o total geral, fora da carga
COL_CODIGO = 2      # B
COL_DESCRICAO = 3   # C
COL_UNIDADE = 5     # E
COL_PRECO = 6       # F
COL_QTD_PREVISTA = 7  # G
COL_MEDICOES = list(range(9, 19))  # I..R, 1ª a 10ª medição

# Linha 20: DOPE repete o código 02.02 da linha 19 na planilha. Decisão do Tiago: entra
# com 02.02.01, renomeada ANTES de montar a hierarquia (pai por maior prefixo).
LINHA_DOPE = 20
CODIGO_DOPE_NOVO = '02.02.01'

GRUPOS = ['01', '02', '03', '04', '05', '06', '07', '08']
DUAS_CASAS = Decimal('0.01')

MEDICOES = [
    {'numero': 1, 'periodo_inicio': '2025-11-01', 'periodo_fim': '2025-11-30'},
    {'numero': 2, 'periodo_inicio': '2025-12-01', 'periodo_fim': '2025-12-31'},
    {'numero': 3, 'periodo_inicio': '2026-01-01', 'periodo_fim': '2026-01-31'},
    {'numero': 4, 'periodo_inicio': '2026-02-01', 'periodo_fim': '2026-02-28'},
    {'numero': 5, 'periodo_inicio': '2026-03-01', 'periodo_fim': '2026-03-31'},
    {'numero': 6, 'periodo_inicio': '2026-04-01', 'periodo_fim': '2026-04-30'},
    {'numero': 7, 'periodo_inicio': '2026-05-01', 'periodo_fim': '2026-05-31'},
    {'numero': 8, 'periodo_inicio': '2026-06-01', 'periodo_fim': '2026-06-30'},
    {'numero': 9, 'periodo_inicio': '2026-07-01', 'periodo_fim': '2026-07-31'},
    {'numero': 10, 'periodo_inicio': '2026-08-01', 'periodo_fim': '2026-08-31'},
]

# Dados do contrato: vault business/lote09-br364, fonte o PDF do contrato SEI 22563019.
# Não vêm da planilha (D3/D4 do "Fonte e decisões" do plano). Só os campos decididos.
CONTRATO = {
    'codigo': 'L09-BR364',
    'nome_obra': 'BR-364/AC Lote 09 (manutenção)',
    'local': 'BR-364/AC, km 620,90 a 682,90',
    'objeto': 'Serviços de manutenção rodoviária (conservação/recuperação) na BR-364/AC',
    'numero_contrato': '00615/2025',
    'contratante_nome': 'DNIT - Superintendência Regional do Acre',
    'contratante_tipo': 'federal',
    'valor_inicial': '243927498.02',
    'data_assinatura': '2025-10-01',
    'prazo_meses': 39,
    'inicio_prazo': 'assinatura',
    'dia_inicio_periodo': 1,
    'tipo_localizacao': 'rodovia',
    'regra_arredondamento': 'sem_arredondar',
    'status': 'ativo',
    'observacoes': (
        'Carga inicial da planilha oficial v12 (Medicao_Teste_3_ATUALIZADA_v12_NOVO.xlsx). '
        'Diferença de R$ 14,53 entre o valor do contrato (R$ 243.927.498,02) e o previsto '
        'da planilha (D4, R$ 243.927.483,49). Linha 20 (DOPE) trocada de 02.02 para '
        '02.02.01: no xlsx repete o código da linha 19 (Usinagem de concreto asfáltico).'
    ),
}


def sha256_arquivo(caminho):
    with open(caminho, 'rb') as f:
        return hashlib.sha256(f.read()).hexdigest()


def numero_texto(valor):
    """Texto exato do número da célula, no mesmo formato do importador do app
    (src/modules/medicao/planilha/leitor.ts, numeroParaTexto: JS String(n)). Um float
    inteiro vira "36", não "36.0"; o não inteiro mantém o repr do Python (mesmos dígitos
    de ida e volta que o JS produz, nas grandezas desta planilha). Recusa notação
    científica."""
    if valor is None:
        return None
    if isinstance(valor, bool) or isinstance(valor, str):
        raise TypeError(f'valor numérico inesperado (tipo {type(valor).__name__}): {valor!r}')
    numero = float(valor)
    texto = repr(numero)
    if 'e' in texto or 'E' in texto:
        raise ValueError(f'número em notação científica, recusado: {valor!r} -> {texto}')
    if numero.is_integer():
        return str(int(numero))
    return texto


def eh_titulo(codigo, preco_cel, qtd_cel):
    """Título: código de 2 dígitos (grupo) ou preço e quantidade prevista vazios."""
    return bool(re.fullmatch(r'\d{2}', codigo)) or (preco_cel is None and qtd_cel is None)


def _segmentos(codigo):
    return tuple(codigo.split('.'))


def atribuir_pai(linhas):
    """Pai de cada linha = a linha anterior de maior prefixo do código (pilha de códigos)."""
    pilha = []  # [(ordem, segmentos)], do mais externo ao mais interno
    for linha in linhas:
        segs = _segmentos(linha['codigo'])
        while pilha and not (len(pilha[-1][1]) < len(segs) and segs[:len(pilha[-1][1])] == pilha[-1][1]):
            pilha.pop()
        linha['pai_ordem'] = pilha[-1][0] if pilha else None
        pilha.append((linha['ordem'], segs))


def carregar_aba(caminho):
    hash_real = sha256_arquivo(caminho)
    if hash_real != SHA256_ESPERADO:
        raise ValueError(
            'hash do xlsx não confere: esperado {}, veio {} ({}). Recusado: use a cópia '
            'oficial do Lote 09 combinada com o Tiago.'.format(SHA256_ESPERADO, hash_real, caminho)
        )
    livro = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    return livro[ABA]


def extrair_linhas(aba):
    linhas = []
    for r in range(PRIMEIRA_LINHA, ULTIMA_LINHA + 1):
        codigo_cel = aba.cell(row=r, column=COL_CODIGO).value
        if codigo_cel is None:
            raise ValueError(f'linha {r}: sem código (linha vazia inesperada no meio da planilha)')
        codigo = str(codigo_cel).strip()
        if r == LINHA_DOPE:
            codigo = CODIGO_DOPE_NOVO

        descricao_cel = aba.cell(row=r, column=COL_DESCRICAO).value
        descricao = str(descricao_cel).strip() if descricao_cel is not None else ''
        if not descricao:
            raise ValueError(f'linha {r} ({codigo}): sem descrição')

        unidade_cel = aba.cell(row=r, column=COL_UNIDADE).value
        unidade = str(unidade_cel).strip() if unidade_cel is not None else None

        preco_cel = aba.cell(row=r, column=COL_PRECO).value
        qtd_cel = aba.cell(row=r, column=COL_QTD_PREVISTA).value

        if eh_titulo(codigo, preco_cel, qtd_cel):
            tipo = 'titulo'
            preco_texto = None
            qtd_texto = None
        else:
            tipo = 'servico'
            if preco_cel is None:
                raise ValueError(f'linha {r} ({codigo}): serviço sem preço')
            preco_texto = numero_texto(preco_cel)
            # Vazio vira "0" em serviço (campo vazio na importação vira zero).
            qtd_texto = numero_texto(qtd_cel) if qtd_cel is not None else '0'

        linhas.append({
            'ordem': len(linhas) + 1,
            'linha_origem': r,
            'codigo': codigo,
            'descricao': descricao,
            'unidade': unidade,
            'tipo': tipo,
            'preco_unitario': preco_texto,
            'quantidade_prevista': qtd_texto,
        })

    atribuir_pai(linhas)
    return linhas


def extrair_quantidades(aba, linhas):
    quantidades = []
    for linha in linhas:
        if linha['tipo'] != 'servico':
            continue
        r = linha['linha_origem']
        for numero_medicao, col in enumerate(COL_MEDICOES, start=1):
            valor_cel = aba.cell(row=r, column=col).value
            # Vazio vira "0" em serviço, mesma regra da quantidade prevista.
            texto = numero_texto(valor_cel) if valor_cel is not None else '0'
            quantidades.append({
                'numero_medicao': numero_medicao,
                'ordem': linha['ordem'],
                'quantidade': texto,
            })
    return quantidades


def _r2(valor):
    return valor.quantize(DUAS_CASAS, rounding=ROUND_HALF_UP)


def _fmt(valor):
    return format(valor, 'f')


def calcular_esperado(linhas, quantidades):
    """Regra sem_arredondar: soma exata (Decimal) por item e por grupo; só arredonda em
    2 casas na saída de cada grupo e do total. O total sai do arredondamento da soma
    exata de TODOS os serviços, não da soma dos grupos já arredondados (por isso o total
    pode diferir em 1 centavo da soma dos grupos: decisão registrada no plano)."""
    servicos = [l for l in linhas if l['tipo'] == 'servico']
    titulos = [l for l in linhas if l['tipo'] == 'titulo']

    qtd_por_chave = {(q['ordem'], q['numero_medicao']): Decimal(q['quantidade']) for q in quantidades}

    grupos_prev = {g: Decimal('0') for g in GRUPOS}
    grupos_acum = {g: Decimal('0') for g in GRUPOS}
    grupos_dec = {g: Decimal('0') for g in GRUPOS}
    total_prev = Decimal('0')
    total_acum = Decimal('0')
    total_dec = Decimal('0')

    precos = {}
    qtds_previstas = {}
    previstos = {}
    qtds_acumuladas = {}
    ajustes = 0

    for linha in servicos:
        ordem = linha['ordem']
        grupo = linha['codigo'].split('.')[0]
        preco = Decimal(linha['preco_unitario'])
        qtd_prevista = Decimal(linha['quantidade_prevista'])

        precos[ordem] = linha['preco_unitario']
        qtds_previstas[ordem] = linha['quantidade_prevista']

        valor_previsto = preco * qtd_prevista
        previstos[ordem] = _fmt(valor_previsto)

        medidas = [qtd_por_chave[(ordem, n)] for n in range(1, 11)]
        acumulada = sum(medidas, Decimal('0'))
        qtds_acumuladas[ordem] = _fmt(acumulada)
        decima_qtd = medidas[9]

        valor_acumulado = preco * acumulada
        valor_decimo = preco * decima_qtd

        grupos_prev[grupo] += valor_previsto
        grupos_acum[grupo] += valor_acumulado
        grupos_dec[grupo] += valor_decimo
        total_prev += valor_previsto
        total_acum += valor_acumulado
        total_dec += valor_decimo

        for qtd in medidas:
            if qtd != 0:
                ajustes += 1

    grupos_saida = {
        g: {
            'previsto': _fmt(_r2(grupos_prev[g])),
            'acumulado': _fmt(_r2(grupos_acum[g])),
            'decima': _fmt(_r2(grupos_dec[g])),
        }
        for g in GRUPOS
    }

    total_previsto = _r2(total_prev)
    total_acumulado = _r2(total_acum)
    total_decima = _r2(total_dec)
    saldo = total_previsto - total_acumulado

    return {
        'linhas': len(linhas),
        'titulos': len(titulos),
        'servicos': len(servicos),
        'medicoes': len(MEDICOES),
        'ajustes': ajustes,
        'grupos': grupos_saida,
        'total': {
            'previsto': _fmt(total_previsto),
            'acumulado': _fmt(total_acumulado),
            'decima': _fmt(total_decima),
            'saldo': _fmt(saldo),
        },
        'precos': precos,
        'qtds_previstas': qtds_previstas,
        'previstos': previstos,
        'qtds_acumuladas': qtds_acumuladas,
    }


def gerar(caminho):
    aba = carregar_aba(caminho)
    linhas = extrair_linhas(aba)
    quantidades = extrair_quantidades(aba, linhas)
    esperado = calcular_esperado(linhas, quantidades)
    staging = {
        'contrato': CONTRATO,
        'linhas': linhas,
        'medicoes': MEDICOES,
        'quantidades': quantidades,
    }
    return staging, esperado


def main():
    caminho = sys.argv[1] if len(sys.argv) > 1 else XLSX_PADRAO
    staging, esperado = gerar(caminho)

    os.makedirs(RETRATO, exist_ok=True)
    with open(os.path.join(RETRATO, 'staging_l09.json'), 'w', encoding='utf-8') as f:
        json.dump(staging, f, ensure_ascii=False, indent=1)
    with open(os.path.join(RETRATO, 'esperado_l09.json'), 'w', encoding='utf-8') as f:
        json.dump(esperado, f, ensure_ascii=False, indent=1)

    print('linhas', esperado['linhas'], 'titulos', esperado['titulos'], 'servicos', esperado['servicos'])
    print('medicoes', esperado['medicoes'], 'ajustes', esperado['ajustes'])
    print('total', esperado['total'])


if __name__ == '__main__':
    main()
