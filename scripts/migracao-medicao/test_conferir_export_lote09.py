"""Testes do conferidor do export do boletim (Task 8, Fase 3 de Medição de Contratos).

Monta com openpyxl, em tempfile, dois arquivos mínimos (3 linhas + total), um no layout da
planilha oficial (aba "Planilha de Medição", cabeçalho na 14, dados a partir da 15) e outro no
layout do export do app (aba "Boletim", marca nas linhas 1 a 5, contexto na 6, cabeçalho na 7,
dados a partir da 8, linha "Total:"), e confere o comportamento do conferidor. Não usa a planilha
oficial de verdade: roda em qualquer máquina com openpyxl. Sem openpyxl, os testes são pulados.

Uso:
  python3 -m unittest scripts/migracao-medicao/test_conferir_export_lote09.py -v
"""
import os
import shutil
import sys
import tempfile
import unittest
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

try:
    import openpyxl
    ERRO_IMPORTACAO = None
except ImportError as erro:  # ex.: CI sem openpyxl
    openpyxl = None
    ERRO_IMPORTACAO = erro

if openpyxl is not None:
    # Só o openpyxl ausente pula; o conferidor ausente ou quebrado tem de falhar.
    import conferir_export_lote09 as ce
else:
    ce = None

MOTIVO_PULO = (
    f'openpyxl ausente: {ERRO_IMPORTACAO}' if ERRO_IMPORTACAO is not None else None
)

MEDICOES = 10

# Três linhas: um título, dois serviços. Valores escolhidos para o TRUNC da planilha oficial
# diferir da conta direta do módulo no serviço 01.01 (H - AT = 100.4567 - 10.1234 = 90.3333:
# TRUNC(…, 3) = 90.333; round(H, 2) - round(AT, 2) = 100.46 - 10.12 = 90.34).
LINHAS = [
    # código, descrição, unidade, preço, qtd prevista, previsto, qtds 1..10, valor 10ª, acumulado
    {
        'codigo': '01', 'descricao': 'TÍTULO', 'unidade': None, 'preco': None, 'qtd': None,
        'previsto': 200.4567, 'qtds': [None] * MEDICOES, 'decima': 1.005, 'acumulado': 20.1234,
    },
    {
        'codigo': '01.01', 'descricao': 'Serviço um', 'unidade': 'm³', 'preco': 21154.63583333333,
        'qtd': 36, 'previsto': 100.4567, 'qtds': [1, 0.6, None, None, None, None, None, None, 0, 0.749996],
        'decima': 1.005, 'acumulado': 10.1234,
    },
    {
        'codigo': '01.02', 'descricao': 'Serviço dois', 'unidade': 'un', 'preco': 102.34700000000001,
        'qtd': 1, 'previsto': 100, 'qtds': [None, 2, None, None, None, None, None, None, None, None],
        'decima': 0, 'acumulado': 10,
    },
]
TOTAL = {'previsto': 200.4567, 'decima': 1.005, 'acumulado': 20.1234}


def _r2(v):
    return float(ce.round_half_up(Decimal(repr(v))))


def _saldo_trunc(previsto, acumulado):
    """TRUNC(H - AT, 3) como a planilha oficial."""
    d = Decimal(repr(previsto)) - Decimal(repr(acumulado))
    return float(d.quantize(Decimal('0.001'), rounding='ROUND_DOWN'))


def _saldo_modulo(previsto, acumulado):
    return float(ce.round_half_up(Decimal(repr(previsto))) - ce.round_half_up(Decimal(repr(acumulado))))


def montar_oficial(caminho, linhas=LINHAS, total=TOTAL):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ce.ABA_OFICIAL
    h = ce.LINHA_CABECALHO_OFICIAL
    ws.cell(h, 2, 'ITEM')
    ws.cell(h, 5, 'Unid.')
    for k in range(36):
        ws.cell(h, 9 + k, f'{k + 1}ª Medição')
    ws.cell(h, 45, f'Valor (R$)\nExecutado\nNa {MEDICOES}ª Medição')
    r = ce.LINHA_CABECALHO_OFICIAL + 1
    for l in linhas:
        ws.cell(r, 2, l['codigo'])
        ws.cell(r, 3, l['descricao'])
        ws.cell(r, 4, 'memória de cálculo')
        ws.cell(r, 5, l['unidade'])
        ws.cell(r, 6, l['preco'])
        ws.cell(r, 7, l['qtd'])
        ws.cell(r, 8, l['previsto'])
        for i, q in enumerate(l['qtds']):
            ws.cell(r, 9 + i, q)
        for c in range(19, 45):  # S..AR: 11ª..36ª, zero nos serviços como na oficial
            ws.cell(r, c, 0 if l['preco'] is not None else None)
        ws.cell(r, 45, l['decima'])
        ws.cell(r, 46, l['acumulado'])
        ws.cell(r, 47, l['acumulado'] / l['previsto'])
        ws.cell(r, 48, _saldo_trunc(l['previsto'], l['acumulado']))
        ws.cell(r, 49, 1 - l['acumulado'] / l['previsto'])
        r += 1
    ws.cell(r, 1, 267)  # como na oficial: A tem um número, B vazia, rótulo em G
    ws.cell(r, 7, 'Total:')
    ws.cell(r, 8, total['previsto'])
    ws.cell(r, 12, 'Total--->')
    ws.cell(r, 45, total['decima'])
    ws.cell(r, 46, total['acumulado'])
    ws.cell(r, 47, total['acumulado'] / total['previsto'])
    ws.cell(r, 48, sum(_saldo_trunc(l['previsto'], l['acumulado']) for l in linhas[1:]))
    ws.cell(r, 49, 1 - total['acumulado'] / total['previsto'])
    wb.save(caminho)


CABECALHO_EXPORT = (
    ['Item', 'Discriminação', 'Unid.', 'Preço Unitário', 'Quantidade Prevista Total',
     'Valor (R$) Previsto Total']
    + [f'{n}ª Medição' for n in range(1, MEDICOES + 1)]
    + [f'Valor (R$) Executado na {MEDICOES}ª Medição', 'Valor (R$) Executado Acumulado',
       'Porcentagem Executada (%)', 'Saldo a Medir (R$)', 'Porcentagem a Medir (%)']
)


def montar_exportado(caminho, linhas=LINHAS, total=TOTAL, mexer=None):
    """Layout do export do app, valores pelas regras do módulo. `mexer(ws)` altera à mão."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Boletim'
    ws.cell(2, 1, 'EMT Construtora')
    ws.cell(6, 1, 'Contrato TESTE · Até a 10ª medição')
    for c, texto in enumerate(CABECALHO_EXPORT, start=1):
        ws.cell(7, c, texto)
    r = 8
    for l in linhas:
        prev = _r2(l['previsto'])
        acum = _r2(l['acumulado'])
        saldo = _saldo_modulo(l['previsto'], l['acumulado'])
        valores = [l['codigo'], l['descricao'], l['unidade'], l['preco'],
                   0 if (l['qtd'] is None and l['preco'] is not None) else l['qtd'], prev]
        valores += [q if q not in (0,) else None for q in l['qtds']]  # 0 da oficial = chave ausente
        valores += [_r2(l['decima']), acum, acum / prev, saldo, saldo / prev]
        for c, v in enumerate(valores, start=1):
            ws.cell(r, c, v)
        r += 1
    prev = _r2(total['previsto'])
    acum = _r2(total['acumulado'])
    saldo = _saldo_modulo(total['previsto'], total['acumulado'])
    ws.cell(r, 1, 'Total:')
    ws.cell(r, 6, prev)
    ws.cell(r, 6 + MEDICOES + 1, _r2(total['decima']))
    ws.cell(r, 6 + MEDICOES + 2, acum)
    ws.cell(r, 6 + MEDICOES + 3, acum / prev)
    ws.cell(r, 6 + MEDICOES + 4, saldo)
    ws.cell(r, 6 + MEDICOES + 5, saldo / prev)
    if mexer:
        mexer(ws)
    wb.save(caminho)


@unittest.skipIf(MOTIVO_PULO is not None, MOTIVO_PULO or '')
class TestConferirExport(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.oficial = os.path.join(self.dir, 'oficial.xlsx')
        self.exportado = os.path.join(self.dir, 'exportado.xlsx')

    def tearDown(self):
        shutil.rmtree(self.dir)

    def conferir(self):
        return ce.conferir(self.oficial, self.exportado, sha256_esperado=None)

    def test_iguais_zero_diferencas_nao_explicadas(self):
        montar_oficial(self.oficial)
        montar_exportado(self.exportado)
        res = self.conferir()
        self.assertEqual(res.linhas_casadas, 3)
        self.assertEqual(res.nao_explicadas, [], ce.formatar_relatorio(res))
        self.assertGreater(res.total_comparadas(), 0)

    def test_um_centavo_a_mais_no_acumulado_acusa_linha_e_coluna(self):
        montar_oficial(self.oficial)

        def mexer(ws):
            # 01.02 está na linha 10 do export; acumulado = coluna R (18)
            ws.cell(10, 18, ws.cell(10, 18).value + 0.01)
        montar_exportado(self.exportado, mexer=mexer)
        res = self.conferir()
        self.assertEqual(len(res.nao_explicadas), 1, ce.formatar_relatorio(res))
        d = res.nao_explicadas[0]
        self.assertEqual(d.linha_oficial, 17)
        self.assertEqual(d.linha_exportada, 10)
        self.assertEqual(d.coluna_oficial, 'AT')
        self.assertEqual(d.coluna_exportada, 'R')
        self.assertEqual(d.codigo, '01.02')
        texto = ce.formatar_relatorio(res)
        self.assertIn('AT', texto)
        self.assertIn('17', texto)

    def test_saldo_direto_contra_trunc_e_explicado(self):
        montar_oficial(self.oficial)
        montar_exportado(self.exportado)
        res = self.conferir()
        self.assertEqual(res.nao_explicadas, [], ce.formatar_relatorio(res))
        saldo = [d for d in res.explicadas if d.coluna_oficial == 'AV']
        # título 01 (soma dos TRUNC), serviço 01.01 e total; 01.02 fecha nas duas contas
        self.assertEqual([d.codigo for d in saldo], ['01', '01.01', 'Total:'],
                         ce.formatar_relatorio(res))
        saldo_linhas = [d for d in saldo if d.codigo == '01.01']
        self.assertEqual(saldo_linhas[0].codigo, '01.01')
        self.assertEqual(saldo_linhas[0].oficial, Decimal('90.333'))
        self.assertEqual(saldo_linhas[0].exportado, Decimal('90.34'))

    def test_saldo_que_nao_segue_a_regra_do_modulo_acusa(self):
        montar_oficial(self.oficial)

        def mexer(ws):
            ws.cell(9, 20, 90.333)  # 01.01 com o TRUNC da oficial: não é a regra do módulo
        montar_exportado(self.exportado, mexer=mexer)
        res = self.conferir()
        self.assertEqual([(d.linha_oficial, d.coluna_oficial) for d in res.nao_explicadas],
                         [(16, 'AV')], ce.formatar_relatorio(res))

    def test_quantidade_com_outro_double_acusa(self):
        montar_oficial(self.oficial)

        def mexer(ws):
            ws.cell(9, 7 + 9, 0.749997)  # 10ª de 01.01 (coluna P)
        montar_exportado(self.exportado, mexer=mexer)
        res = self.conferir()
        self.assertEqual([(d.linha_oficial, d.coluna_oficial, d.coluna_exportada)
                          for d in res.nao_explicadas], [(16, 'R', 'P')])

    def test_medicao_alem_da_decima_na_oficial_acusa(self):
        linhas = [dict(l) for l in LINHAS]
        montar_oficial(self.oficial, linhas)
        wb = openpyxl.load_workbook(self.oficial)
        wb[ce.ABA_OFICIAL].cell(17, 19, 1.5)  # S17 = 11ª de 01.02
        wb.save(self.oficial)
        montar_exportado(self.exportado)
        res = self.conferir()
        self.assertEqual([(d.linha_oficial, d.coluna_oficial) for d in res.nao_explicadas],
                         [(17, 'S')])

    def test_linha_a_mais_no_export_acusa(self):
        montar_oficial(self.oficial)
        extra = LINHAS + [dict(LINHAS[2], codigo='01.03')]
        montar_exportado(self.exportado, linhas=extra)
        res = self.conferir()
        self.assertTrue(any('linhas' in d.motivo for d in res.nao_explicadas),
                        ce.formatar_relatorio(res))

    def test_percentual_fora_da_tolerancia_acusa(self):
        montar_oficial(self.oficial)

        def mexer(ws):
            ws.cell(10, 19, ws.cell(10, 19).value + 0.0001)  # % executada de 01.02
        montar_exportado(self.exportado, mexer=mexer)
        res = self.conferir()
        self.assertEqual([(d.linha_oficial, d.coluna_oficial) for d in res.nao_explicadas],
                         [(17, 'AU')])

    def test_recusa_hash_diferente(self):
        montar_oficial(self.oficial)
        montar_exportado(self.exportado)
        with self.assertRaises(SystemExit):
            ce.conferir(self.oficial, self.exportado)


if __name__ == '__main__':
    unittest.main()
