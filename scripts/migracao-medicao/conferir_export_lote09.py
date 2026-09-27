"""Conferência célula a célula do export do boletim contra a planilha oficial do Lote 09.

Task 8 da Fase 3 de Medição de Contratos. Compara o xlsx que o app exporta (aba "Boletim",
Task 7) com a planilha oficial do DNIT (`Medicao_Teste_3_ATUALIZADA_v12_NOVO.xlsx`, aba
"Planilha de Medição", cabeçalho na linha 14, dados a partir da 15, linha "Total:" com o
rótulo na coluna G). A
planilha oficial só é aceita com o sha256 combinado com o Tiago.

Uso:
  python3 scripts/migracao-medicao/conferir_export_lote09.py <oficial.xlsx> <exportado.xlsx>
      [--relatorio caminho.txt]

Imprime o relatório, grava em `_retrato/conferencia_export.txt` (ou em `--relatorio`) e sai
com 1 se houver qualquer diferença não explicada.

Regras (linha oficial r casa com a linha exportada de mesma ordem; oficial lido com
`data_only=True`; números sempre por `Decimal(repr(v))`):
- B código: texto igual. C descrição e E unidade: iguais depois de aparadas.
- F preço, G quantidade prevista, 1ª..Nª medição: o mesmo double (vazio = 0 = vazio).
- Da (N+1)ª à 36ª medição da oficial: vazio ou zero.
- H previsto, valor na Nª, AT acumulado: round_half_up(oficial, 2) == exportado.
- AU e AW (%): |oficial - exportado| < 0.00005.
- AV saldo: exportado == round(H, 2) - round(AT, 2) da oficial (regra do Tiago); a
  diferença contra o AV oficial (TRUNC(H - AT, 3) nos serviços, soma nos títulos) é
  explicada e contada à parte.

Diferenças explicadas (decisões do Tiago, 26/09/2026, e regra da fase):
- `codigo_dope`: linha 20 da oficial `02.02` virou `02.02.01` no módulo.
- `saldo_trunc`: saldo pela conta direta, não pelo TRUNC da oficial.
- `qtd_prevista_vazia`: serviço com quantidade prevista vazia na oficial (03.16.x) vale 0.
- `unidade_aparada`: unidade com espaço sobrando na oficial (`'un '`) gravada aparada.
- `pct_previsto_zero`: % com previsto zero sai `#DIV/0!` na oficial e vazio no export
  (contexto da Fase 3: "% nulos quando o previsto é zero").
- `subtitulo_em_branco`: subtítulo (título abaixo do nível 1) com a célula vazia na oficial.
  Só é explicada se o valor exportado for o recalculado da própria oficial: H, AS e AT =
  round_half_up(soma das linhas de serviço da subárvore, 2), cada linha com preço contada uma
  vez, títulos fora da soma; AV = round(H) - round(AT) desses; AU e AW a menos de 0.00005 da
  razão desses (nulo com previsto zero). Subárvore = linhas seguintes cujo código começa com
  o do subtítulo + "." (linha 20 como 02.02.01, a mesma hierarquia do módulo). Célula da
  oficial preenchida segue a regra normal, salvo o caso abaixo.
- `formula_servico_em_titulo`: célula de título na oficial com fórmula de serviço que
  referencia `F<linha>`/`$F<linha>` da própria linha, com esse F vazio (o AS102 = 0 do 02.10).
  Só é explicada se o título for título no export (sem preço e sem unidade) e o valor
  exportado passar na mesma verificação do subtotal recalculado da subárvore. Fórmula lida
  numa segunda carga com `data_only=False`.
Qualquer outra diferença é não explicada e aparece no relatório: o conferidor não corrige nada.
"""
import argparse
import hashlib
import os
import re
import sys
import warnings
from collections import OrderedDict
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

import openpyxl
from openpyxl.utils import get_column_letter

# Avisos inofensivos do openpyxl: a oficial tem validação de dados em extensão e o export
# tem a logo (desenho). Só esses dois são calados; qualquer outro aviso continua aparecendo.
warnings.filterwarnings('ignore', message=r'Data Validation extension', category=UserWarning)
warnings.filterwarnings('ignore', message=r'DrawingML support is incomplete', category=UserWarning)

SHA256_ESPERADO = '2a29e7473cca3e057a700a91c58d8d992e89ea42e2c74a094d1f3e59121fcb0b'
ABA_OFICIAL = 'Planilha de Medição'
ABA_EXPORT = 'Boletim'
LINHA_CABECALHO_OFICIAL = 14
# Na oficial de verdade (conferida pelo sha256): 265 linhas, total na 280.
LINHAS_ESPERADAS = 265
LINHA_TOTAL_ESPERADA = 280

# Colunas fixas da oficial.
COL_CODIGO, COL_DESCRICAO, COL_UNIDADE = 2, 3, 5          # B, C, E
COL_PRECO, COL_QTD, COL_PREVISTO = 6, 7, 8                # F, G, H
COL_PRIMEIRA_MEDICAO, MEDICOES_OFICIAL = 9, 36            # I .. AR
COL_VALOR_N, COL_ACUMULADO = 45, 46                       # AS, AT
COL_PCT_EXEC, COL_SALDO, COL_PCT_MEDIR = 47, 48, 49       # AU, AV, AW

TOLERANCIA_PCT = Decimal('0.00005')
DOIS = Decimal('0.01')

EXPLICACOES = OrderedDict([
    ('codigo_dope', 'linha 20 da oficial: código 02.02 virou 02.02.01 no módulo (Tiago, 26/09)'),
    ('saldo_trunc', 'saldo = round(H,2) - round(AT,2) no módulo; oficial usa TRUNC(H-AT,3) '
                    'nos serviços e soma nos títulos (Tiago, 26/09)'),
    ('qtd_prevista_vazia', 'quantidade prevista vazia na oficial vale 0 no módulo (03.16.x, Tiago, 26/09)'),
    ('unidade_aparada', 'unidade com espaço sobrando na oficial gravada aparada (Tiago, 26/09)'),
    ('pct_previsto_zero', '% com previsto zero: #DIV/0! na oficial, vazio no módulo '
                          '(regra da Fase 3: % nulo quando o previsto é zero)'),
    ('subtitulo_em_branco', 'subtítulo em branco na oficial; o módulo mostra o subtotal da subárvore, '
                            'conferido contra a soma dos serviços da própria oficial (controlador, 27/09)'),
    ('formula_servico_em_titulo', 'fórmula de serviço em título na oficial (sobre o F vazio da própria linha); '
                                  'o módulo mostra o subtotal da subárvore, conferido contra a soma dos '
                                  'serviços da própria oficial (controlador, 27/09)'),
])


def sha256_arquivo(caminho):
    with open(caminho, 'rb') as f:
        return hashlib.sha256(f.read()).hexdigest()


def round_half_up(valor, casas=DOIS):
    return valor.quantize(casas, rounding=ROUND_HALF_UP)


def vazio(v):
    return v is None or (isinstance(v, str) and v.strip() == '')


def numero(v):
    """Decimal(repr(v)) de uma célula numérica; None se vazia; levanta ValueError se texto."""
    if vazio(v):
        return None
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        raise ValueError(f'célula não numérica: {v!r}')
    try:
        return Decimal(repr(v))
    except InvalidOperation as erro:  # pragma: no cover (repr de float sempre é Decimal válido)
        raise ValueError(f'célula não numérica: {v!r}') from erro


def _normal(texto):
    return re.sub(r'\s+', ' ', str(texto)).strip() if texto is not None else ''


@dataclass
class Diferenca:
    linha_oficial: object
    linha_exportada: object
    codigo: str
    campo: str
    coluna_oficial: str
    coluna_exportada: str
    oficial: object
    exportado: object
    motivo: str
    explicacao: str = ''
    tipo: str = ''          # 'serviço', 'título' ou 'total' (pela oficial)


class Resultado:
    def __init__(self):
        self.contagem = OrderedDict()   # campo -> {comparadas, ok, explicadas, nao_explicadas}
        self.explicadas = []
        self.nao_explicadas = []
        self.linhas_oficial = 0
        self.linhas_exportadas = 0
        self.linhas_casadas = 0
        self.linha_total_oficial = None
        self.linha_total_exportada = None
        self.medicoes = None
        self.arquivo_oficial = ''
        self.arquivo_exportado = ''
        self.sha256_oficial = ''
        self.totais = []                # (campo, oficial, exportado)

    def _conta(self, campo):
        return self.contagem.setdefault(
            campo, {'comparadas': 0, 'ok': 0, 'explicadas': 0, 'nao_explicadas': 0})

    def ok(self, campo):
        c = self._conta(campo)
        c['comparadas'] += 1
        c['ok'] += 1

    def explicada(self, diferenca, chave):
        c = self._conta(diferenca.campo)
        c['comparadas'] += 1
        c['explicadas'] += 1
        diferenca.explicacao = chave
        self.explicadas.append(diferenca)

    def nao_explicada(self, diferenca, contar=True):
        if contar:
            c = self._conta(diferenca.campo)
            c['comparadas'] += 1
            c['nao_explicadas'] += 1
        self.nao_explicadas.append(diferenca)

    def total_comparadas(self):
        return sum(c['comparadas'] for c in self.contagem.values())


# ---------------------------------------------------------------- leitura dos dois arquivos

def _abrir_oficial(caminho, sha256_esperado):
    if sha256_esperado is not None:
        real = sha256_arquivo(caminho)
        if real != sha256_esperado:
            raise SystemExit(
                f'planilha oficial recusada: sha256 {real} não é o combinado ({sha256_esperado})')
    wb = openpyxl.load_workbook(caminho, data_only=True)
    if ABA_OFICIAL not in wb.sheetnames:
        raise SystemExit(f'planilha oficial sem a aba "{ABA_OFICIAL}"')
    ws = wb[ABA_OFICIAL]

    h = LINHA_CABECALHO_OFICIAL
    esperado = {COL_CODIGO: 'ITEM', COL_UNIDADE: 'Unid.'}
    for k in range(MEDICOES_OFICIAL):
        esperado[COL_PRIMEIRA_MEDICAO + k] = f'{k + 1}ª Medição'
    for col, texto in esperado.items():
        achado = _normal(ws.cell(h, col).value)
        if achado != texto:
            raise SystemExit(f'layout da oficial mudou: {get_column_letter(col)}{h} = {achado!r}, '
                             f'esperado {texto!r}')
    m = re.search(r'Na (\d+)ª Medição', _normal(ws.cell(h, COL_VALOR_N).value), re.IGNORECASE)
    if not m:
        raise SystemExit(f'layout da oficial mudou: {get_column_letter(COL_VALOR_N)}{h} = '
                         f'{ws.cell(h, COL_VALOR_N).value!r}')
    medicoes = int(m.group(1))

    linhas = []
    r = h + 1
    while r <= ws.max_row:
        # Na oficial o rótulo "Total:" fica na coluna G (B vazia); aceita de B a G.
        if any(_normal(ws.cell(r, c).value) == 'Total:' for c in range(COL_CODIGO, COL_QTD + 1)):
            formulas = openpyxl.load_workbook(caminho, data_only=False)[ABA_OFICIAL]
            return ws, formulas, linhas, r, medicoes
        linhas.append(r)
        r += 1
    raise SystemExit('planilha oficial sem a linha "Total:" (colunas B a G)')


def _abrir_exportado(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb[ABA_EXPORT] if ABA_EXPORT in wb.sheetnames else wb.active
    cab = None
    for r in range(1, min(ws.max_row, 40) + 1):
        if _normal(ws.cell(r, 1).value) == 'Item':
            cab = r
            break
    if cab is None:
        raise SystemExit('export sem a linha de cabeçalho (coluna A = "Item")')
    colunas = {}
    medicoes = {}
    for c in range(1, ws.max_column + 1):
        texto = _normal(ws.cell(cab, c).value)
        if not texto:
            continue
        m = re.fullmatch(r'(\d+)ª Medição', texto)
        if m:
            medicoes[int(m.group(1))] = c
        else:
            colunas[texto] = c
    linhas = []
    r = cab + 1
    while r <= ws.max_row:
        if _normal(ws.cell(r, 1).value) == 'Total:':
            return ws, cab, colunas, medicoes, linhas, r
        linhas.append(r)
        r += 1
    raise SystemExit('export sem a linha "Total:" na coluna A')


def _coluna_export(colunas, texto):
    if texto not in colunas:
        raise SystemExit(f'export sem a coluna "{texto}"')
    return colunas[texto]


# ---------------------------------------------------------------- regras

def _diferenca(ctx, campo, col_o, col_e, oficial, exportado, motivo):
    return Diferenca(ctx['linha_oficial'], ctx['linha_exportada'], ctx['codigo'], campo,
                     get_column_letter(col_o) if col_o else '-',
                     get_column_letter(col_e) if col_e else '-', oficial, exportado, motivo,
                     tipo=ctx.get('tipo', ''))


def _texto(res, ctx, campo, col_o, col_e, o, e, aparar, explicacao_aparada=None):
    so = '' if o is None else str(o)
    se = '' if e is None else str(e)
    if aparar:
        if so.strip() == se.strip():
            if explicacao_aparada and so != so.strip() and se == se.strip():
                res.explicada(_diferenca(ctx, campo, col_o, col_e, o, e, 'espaço sobrando na oficial'),
                              explicacao_aparada)
            else:
                res.ok(campo)
            return
    elif so == se:
        res.ok(campo)
        return
    res.nao_explicada(_diferenca(ctx, campo, col_o, col_e, o, e, 'texto diferente'))


def _mesmo_double(res, ctx, campo, col_o, col_e, o, e, servico=False, explicar_vazio=None):
    try:
        do, de = numero(o), numero(e)
    except ValueError as erro:
        res.nao_explicada(_diferenca(ctx, campo, col_o, col_e, o, e, str(erro)))
        return
    if do is None and de is None:
        res.ok(campo)
        return
    if do is None or de is None:
        outro = de if do is None else do
        if outro == 0:
            if explicar_vazio and servico and do is None:
                res.explicada(_diferenca(ctx, campo, col_o, col_e, o, e, 'vazio na oficial, 0 no export'),
                              explicar_vazio)
            else:
                res.ok(campo)
            return
        res.nao_explicada(_diferenca(ctx, campo, col_o, col_e, o, e, 'vazio de um lado, número do outro'))
        return
    if do == de:
        res.ok(campo)
    else:
        res.nao_explicada(_diferenca(ctx, campo, col_o, col_e, o, e, 'double diferente'))


def _dinheiro(res, ctx, campo, col_o, col_e, o, e):
    try:
        do, de = numero(o), numero(e)
    except ValueError as erro:
        res.nao_explicada(_diferenca(ctx, campo, col_o, col_e, o, e, str(erro)))
        return
    do = Decimal(0) if do is None else do
    if de is None:
        if do == 0:
            res.ok(campo)
        else:
            res.nao_explicada(_diferenca(ctx, campo, col_o, col_e, o, e, 'vazio no export'))
        return
    if round_half_up(do) == de:
        res.ok(campo)
    else:
        extra = ''
        sub = ctx.get('subtitulo')
        if sub is not None and not sub['erro']:
            chave = {COL_PREVISTO: 'previsto', COL_VALOR_N: 'valor_n', COL_ACUMULADO: 'acumulado'}.get(col_o)
            if chave:
                extra = (f'; célula preenchida na oficial, fora da regra do subtítulo '
                         f'(Σ serviços da subárvore = {sub[chave]})')
        res.nao_explicada(_diferenca(
            ctx, campo, col_o, col_e, o, e,
            f'round_half_up(oficial, 2) = {round_half_up(do)} diferente do exportado {de}{extra}'))


def _pct(res, ctx, campo, col_o, col_e, o, e, previsto_oficial):
    if isinstance(o, str) and o.strip() == '#DIV/0!' and e is None:
        if previsto_oficial == 0:
            res.explicada(_diferenca(ctx, campo, col_o, col_e, o, e, 'previsto zero'), 'pct_previsto_zero')
            return
    try:
        do, de = numero(o), numero(e)
    except ValueError as erro:
        res.nao_explicada(_diferenca(ctx, campo, col_o, col_e, o, e, str(erro)))
        return
    if do is None and de is None:
        res.ok(campo)
        return
    if do is None or de is None:
        res.nao_explicada(_diferenca(ctx, campo, col_o, col_e, o, e, 'vazio de um lado, número do outro'))
        return
    if abs(do - de) < TOLERANCIA_PCT:
        res.ok(campo)
    else:
        res.nao_explicada(_diferenca(ctx, campo, col_o, col_e, o, e,
                                     f'|oficial - exportado| = {abs(do - de)} >= {TOLERANCIA_PCT}'))


def _saldo(res, ctx, campo, col_o, col_e, o, e, previsto_o, acumulado_o):
    try:
        do, de = numero(o), numero(e)
        dh, dat = numero(previsto_o), numero(acumulado_o)
    except ValueError as erro:
        res.nao_explicada(_diferenca(ctx, campo, col_o, col_e, o, e, str(erro)))
        return
    esperado = round_half_up(dh or Decimal(0)) - round_half_up(dat or Decimal(0))
    if de is None or de != esperado:
        res.nao_explicada(_diferenca(
            ctx, campo, col_o, col_e, o, e,
            f'exportado não é round(H,2) - round(AT,2) da oficial = {esperado}'))
        return
    if (do if do is not None else Decimal(0)) == de:
        res.ok(campo)
        return
    res.explicada(_diferenca(ctx, campo, col_o, col_e, do, de,
                             f'oficial - exportado = {(do or Decimal(0)) - de}'), 'saldo_trunc')


# ---------------------------------------------------------------- conferência

def conferir(caminho_oficial, caminho_exportado, sha256_esperado=SHA256_ESPERADO):
    res = Resultado()
    res.arquivo_oficial = os.path.abspath(caminho_oficial)
    res.arquivo_exportado = os.path.abspath(caminho_exportado)
    res.sha256_oficial = sha256_arquivo(caminho_oficial)

    wo, wf, linhas_o, total_o, n = _abrir_oficial(caminho_oficial, sha256_esperado)
    we, cab_e, colunas, medicoes_e, linhas_e, total_e = _abrir_exportado(caminho_exportado)
    res.medicoes = n
    res.linhas_oficial, res.linhas_exportadas = len(linhas_o), len(linhas_e)
    res.linha_total_oficial, res.linha_total_exportada = total_o, total_e

    geral = {'linha_oficial': '-', 'linha_exportada': '-', 'codigo': '-'}
    if sorted(medicoes_e) != list(range(1, n + 1)):
        res.nao_explicada(_diferenca(geral, 'layout', None, None, f'1ª..{n}ª',
                                     sorted(medicoes_e), 'colunas de medição do export'), contar=False)
    ce = {
        'codigo': _coluna_export(colunas, 'Item'),
        'descricao': _coluna_export(colunas, 'Discriminação'),
        'unidade': _coluna_export(colunas, 'Unid.'),
        'preco': _coluna_export(colunas, 'Preço Unitário'),
        'qtd': _coluna_export(colunas, 'Quantidade Prevista Total'),
        'previsto': _coluna_export(colunas, 'Valor (R$) Previsto Total'),
        'valor_n': _coluna_export(colunas, f'Valor (R$) Executado na {n}ª Medição'),
        'acumulado': _coluna_export(colunas, 'Valor (R$) Executado Acumulado'),
        'pct_exec': _coluna_export(colunas, 'Porcentagem Executada (%)'),
        'saldo': _coluna_export(colunas, 'Saldo a Medir (R$)'),
        'pct_medir': _coluna_export(colunas, 'Porcentagem a Medir (%)'),
    }

    if len(linhas_o) != len(linhas_e):
        res.nao_explicada(_diferenca(geral, 'layout', None, None, len(linhas_o), len(linhas_e),
                                     'número de linhas diferente (oficial x export)'), contar=False)
    if sha256_esperado is not None and (len(linhas_o) != LINHAS_ESPERADAS
                                        or total_o != LINHA_TOTAL_ESPERADA):
        res.nao_explicada(_diferenca(geral, 'layout', None, None,
                                     f'{LINHAS_ESPERADAS} linhas, total na {LINHA_TOTAL_ESPERADA}',
                                     f'{len(linhas_o)} linhas, total na {total_o}',
                                     'número de linhas da oficial fora do combinado'), contar=False)

    vo = lambda r, c: wo.cell(r, c).value  # noqa: E731
    ve = lambda r, c: we.cell(r, c).value  # noqa: E731

    subtotais = _subtotais_subtitulos(wo, linhas_o)

    for ro, rex in zip(linhas_o, linhas_e):
        res.linhas_casadas += 1
        codigo_o = vo(ro, COL_CODIGO)
        servico = _eh_servico(wo, ro)
        sub = subtotais.get(ro)
        titulo_no_export = vazio(ve(rex, ce['preco'])) and vazio(ve(rex, ce['unidade']))
        ctx = {'linha_oficial': ro, 'linha_exportada': rex, 'codigo': str(codigo_o),
               'tipo': 'serviço' if servico else 'título',
               'subtitulo': sub if (sub is not None and sub['subtitulo']) else None,
               'subarvore': sub if titulo_no_export else None,
               'formulas': wf}

        # B código
        e_cod = ve(rex, ce['codigo'])
        if ro == 20 and str(codigo_o) == '02.02' and str(e_cod) == '02.02.01':
            res.explicada(_diferenca(ctx, 'código', COL_CODIGO, ce['codigo'], codigo_o, e_cod,
                                     'código repetido na oficial'), 'codigo_dope')
        else:
            _texto(res, ctx, 'código', COL_CODIGO, ce['codigo'], codigo_o, e_cod, aparar=False)
        _texto(res, ctx, 'descrição', COL_DESCRICAO, ce['descricao'],
               vo(ro, COL_DESCRICAO), ve(rex, ce['descricao']), aparar=True)
        _texto(res, ctx, 'unidade', COL_UNIDADE, ce['unidade'],
               vo(ro, COL_UNIDADE), ve(rex, ce['unidade']), aparar=True,
               explicacao_aparada='unidade_aparada')

        _mesmo_double(res, ctx, 'preço', COL_PRECO, ce['preco'], vo(ro, COL_PRECO), ve(rex, ce['preco']))
        _mesmo_double(res, ctx, 'qtd prevista', COL_QTD, ce['qtd'], vo(ro, COL_QTD), ve(rex, ce['qtd']),
                      servico=servico, explicar_vazio='qtd_prevista_vazia')
        for k in range(1, n + 1):
            col_o = COL_PRIMEIRA_MEDICAO + k - 1
            col_e = medicoes_e.get(k)
            _mesmo_double(res, ctx, f'{k}ª medição', col_o, col_e, vo(ro, col_o),
                          ve(rex, col_e) if col_e else None)
        for k in range(n + 1, MEDICOES_OFICIAL + 1):
            col_o = COL_PRIMEIRA_MEDICAO + k - 1
            v = vo(ro, col_o)
            try:
                d = numero(v)
            except ValueError as erro:
                res.nao_explicada(_diferenca(ctx, f'{n + 1}ª..36ª (oficial)', col_o, None, v, None, str(erro)))
                continue
            if d is None or d == 0:
                res.ok(f'{n + 1}ª..36ª (oficial)')
            else:
                res.nao_explicada(_diferenca(ctx, f'{n + 1}ª..36ª (oficial)', col_o, None, v, None,
                                             f'medição além da {n}ª preenchida na oficial'))

        _conferir_valores(res, ctx, vo, ve, ro, rex, ce)

    # Linha do total
    ctx = {'linha_oficial': total_o, 'linha_exportada': total_e, 'codigo': 'Total:', 'tipo': 'total'}
    _conferir_valores(res, ctx, vo, ve, total_o, total_e, ce, sufixo=' (total)')
    for campo, co, key in [('previsto', COL_PREVISTO, 'previsto'), ('valor na Nª', COL_VALOR_N, 'valor_n'),
                           ('acumulado', COL_ACUMULADO, 'acumulado'), ('% executada', COL_PCT_EXEC, 'pct_exec'),
                           ('saldo', COL_SALDO, 'saldo'), ('% a medir', COL_PCT_MEDIR, 'pct_medir')]:
        res.totais.append((campo, get_column_letter(co), vo(total_o, co),
                           get_column_letter(ce[key]), ve(total_e, ce[key])))
    return res


def _conferir_valores(res, ctx, vo, ve, ro, rex, ce, sufixo=''):
    previsto_o = vo(ro, COL_PREVISTO)
    try:
        dprev = numero(previsto_o)
    except ValueError:
        dprev = None
    sub = ctx.get('subtitulo')
    colunas = [
        ('previsto', COL_PREVISTO, 'previsto'), ('valor na Nª', COL_VALOR_N, 'valor_n'),
        ('acumulado', COL_ACUMULADO, 'acumulado'), ('% executada', COL_PCT_EXEC, 'pct_exec'),
        ('saldo', COL_SALDO, 'saldo'), ('% a medir', COL_PCT_MEDIR, 'pct_medir'),
    ]
    for campo, col_o, chave in colunas:
        o, e = vo(ro, col_o), ve(rex, ce[chave])
        if sub is not None and vazio(o):
            _subtitulo(res, ctx, campo + sufixo, col_o, ce[chave], o, e, sub, chave)
        elif _formula_de_servico_sobre_f_vazio(ctx, ro, col_o, vo):
            _formula_titulo(res, ctx, campo + sufixo, col_o, ce[chave], o, e, chave)
        elif chave in ('previsto', 'valor_n', 'acumulado'):
            _dinheiro(res, ctx, campo + sufixo, col_o, ce[chave], o, e)
        elif chave in ('pct_exec', 'pct_medir'):
            _pct(res, ctx, campo + sufixo, col_o, ce[chave], o, e, dprev)
        else:
            _saldo(res, ctx, campo + sufixo, col_o, ce[chave], o, e, previsto_o, vo(ro, COL_ACUMULADO))


# ---------------------------------------------------------------- subtítulos

def _formula_de_servico_sobre_f_vazio(ctx, ro, col_o, vo):
    """Título na oficial e no export, célula com fórmula que usa o F da própria linha, F vazio."""
    if ctx.get('subarvore') is None or ctx.get('formulas') is None:
        return False
    formula = ctx['formulas'].cell(ro, col_o).value
    if not (isinstance(formula, str) and formula.startswith('=')):
        return False
    if not re.search(rf'(?<![A-Z$])\$?F\$?{ro}(?!\d)', formula):
        return False
    return vazio(vo(ro, COL_PRECO))


def _formula_titulo(res, ctx, campo, col_o, col_e, o, e, chave):
    sub = ctx['subarvore']
    formula = ctx['formulas'].cell(ctx['linha_oficial'], col_o).value
    try:
        do, de = numero(o), numero(e)
    except ValueError as erro:
        res.nao_explicada(_diferenca(ctx, campo, col_o, col_e, o, e, str(erro)))
        return
    if chave in ('pct_exec', 'pct_medir'):
        if do is not None and de is not None and abs(do - de) < TOLERANCIA_PCT:
            res.ok(campo)
            return
        esperado = sub[chave]
        passou = (not sub['erro'] and esperado is not None and de is not None
                  and abs(de - esperado) < TOLERANCIA_PCT)
    else:
        if de is not None and round_half_up(do or Decimal(0)) == de:
            res.ok(campo)
            return
        esperado = sub[chave]
        passou = not sub['erro'] and de is not None and de == esperado
    base = (f'fórmula {formula} sobre F{ctx["linha_oficial"]} vazio; '
            f'Σ de {len(sub["linhas"])} serviço(s) da subárvore na oficial = {esperado}')
    if passou:
        res.explicada(_diferenca(ctx, campo, col_o, col_e, o, e, base), 'formula_servico_em_titulo')
    else:
        res.nao_explicada(_diferenca(ctx, campo, col_o, col_e, o, e, base + ', diferente do exportado'))


def _eh_servico(ws, r):
    return not vazio(ws.cell(r, COL_PRECO).value) and _normal(ws.cell(r, COL_UNIDADE).value) != ''


def _codigo_hierarquia(ws, r):
    """Código na hierarquia do módulo: a linha 20 (02.02 repetido, DOPE) é 02.02.01."""
    codigo = _normal(ws.cell(r, COL_CODIGO).value)
    return '02.02.01' if (r == 20 and codigo == '02.02') else codigo


def _subtotais_subtitulos(ws, linhas):
    """{linha: subtotais} dos subtítulos (títulos abaixo do nível 1), recalculados da oficial.

    Soma exata (Decimal(repr)) de H, AS e AT das linhas de serviço da subárvore (linhas
    seguintes com código começando por código + "."), cada linha com preço uma vez, títulos
    fora; depois round_half_up em 2 casas. Célula não numérica numa linha da soma deixa o
    subtítulo sem subtotal ("erro"), e aí nada dele é explicado.
    """
    codigos = [(r, _codigo_hierarquia(ws, r)) for r in linhas]
    saida = {}
    for i, (r, codigo) in enumerate(codigos):
        if _eh_servico(ws, r):
            continue
        somas = {'H': Decimal(0), 'AS': Decimal(0), 'AT': Decimal(0)}
        erro = None
        linhas_soma = []
        for r2, c2 in codigos[i + 1:]:
            if not c2.startswith(codigo + '.'):
                break
            if not _eh_servico(ws, r2):
                continue
            linhas_soma.append(r2)
            for chave, col in (('H', COL_PREVISTO), ('AS', COL_VALOR_N), ('AT', COL_ACUMULADO)):
                try:
                    somas[chave] += numero(ws.cell(r2, col).value) or Decimal(0)
                except ValueError as e:
                    erro = f'linha {r2} col {get_column_letter(col)}: {e}'
        h, as_, at = (round_half_up(somas[k]) for k in ('H', 'AS', 'AT'))
        saldo = h - at
        saida[r] = {
            'erro': erro, 'linhas': linhas_soma, 'subtitulo': '.' in codigo,
            'previsto': h, 'valor_n': as_, 'acumulado': at, 'saldo': saldo,
            'pct_exec': (at / h) if h != 0 else None,
            'pct_medir': (saldo / h) if h != 0 else None,
        }
    return saida


def _subtitulo(res, ctx, campo, col_o, col_e, o, e, sub, chave):
    if sub['erro']:
        res.nao_explicada(_diferenca(ctx, campo, col_o, col_e, o, e,
                                     f'subtítulo sem subtotal recalculável ({sub["erro"]})'))
        return
    esperado = sub[chave]
    base = f'Σ de {len(sub["linhas"])} serviço(s) da subárvore na oficial'
    try:
        de = numero(e)
    except ValueError as erro:
        res.nao_explicada(_diferenca(ctx, campo, col_o, col_e, o, e, str(erro)))
        return
    if chave in ('pct_exec', 'pct_medir'):
        if esperado is None:
            if de is None:
                res.ok(campo)
            else:
                res.nao_explicada(_diferenca(ctx, campo, col_o, col_e, o, e,
                                             f'{base}: previsto zero, % deveria ser vazio'))
            return
        if de is not None and abs(de - esperado) < TOLERANCIA_PCT:
            res.explicada(_diferenca(ctx, campo, col_o, col_e, o, e,
                                     f'{base}: razão recalculada {esperado:.10f}'), 'subtitulo_em_branco')
        else:
            res.nao_explicada(_diferenca(ctx, campo, col_o, col_e, o, e,
                                         f'{base}: razão recalculada {esperado:.10f}, fora da tolerância'))
        return
    # dinheiro e saldo
    if esperado == 0 and (de is None or de == 0):
        res.ok(campo)  # vazio na oficial e zero no export, como em qualquer linha
        return
    if de is not None and de == esperado:
        res.explicada(_diferenca(ctx, campo, col_o, col_e, o, e, f'{base} = {esperado}'),
                      'subtitulo_em_branco')
    else:
        res.nao_explicada(_diferenca(ctx, campo, col_o, col_e, o, e,
                                     f'{base} = {esperado}, diferente do exportado'))


# ---------------------------------------------------------------- relatório

def _v(x):
    if isinstance(x, float):
        return repr(x)
    return 'vazio' if x is None else str(x)


def formatar_relatorio(res):
    out = []
    p = out.append
    p('Conferência célula a célula: export do boletim x planilha oficial do Lote 09')
    p('')
    p(f'Oficial:   {res.arquivo_oficial}')
    p(f'           sha256 {res.sha256_oficial}')
    p(f'Exportado: {res.arquivo_exportado}')
    p(f'Boletim até a {res.medicoes}ª medição')
    p(f'Linhas: oficial {res.linhas_oficial} (total na {res.linha_total_oficial}), '
      f'exportado {res.linhas_exportadas} (total na {res.linha_total_exportada}); '
      f'casadas {res.linhas_casadas}')
    p('')
    p('Contagem por campo (células comparadas):')
    largura = max([len(c) for c in res.contagem] + [5])
    p(f'  {"campo".ljust(largura)}  comparadas      ok  explicadas  não explicadas')
    soma = {'comparadas': 0, 'ok': 0, 'explicadas': 0, 'nao_explicadas': 0}
    for campo, c in res.contagem.items():
        p(f'  {campo.ljust(largura)}  {c["comparadas"]:>10}  {c["ok"]:>6}  {c["explicadas"]:>10}  '
          f'{c["nao_explicadas"]:>14}')
        for k in soma:
            soma[k] += c[k]
    p(f'  {"TOTAL".ljust(largura)}  {soma["comparadas"]:>10}  {soma["ok"]:>6}  {soma["explicadas"]:>10}  '
      f'{soma["nao_explicadas"]:>14}')
    p('')
    p('Linha do total (oficial x exportado):')
    for campo, co, o, cx, e in res.totais:
        p(f'  {campo:<12} {co:>2} {_v(o):>22}   {cx:>2} {_v(e):>22}')
    p('')
    p(f'Diferenças explicadas: {len(res.explicadas)}')
    for chave, texto in EXPLICACOES.items():
        grupo = [d for d in res.explicadas if d.explicacao == chave]
        if not grupo:
            continue
        p(f'  [{chave}] {texto}: {len(grupo)}')
        if chave == 'saldo_trunc':
            for tipo in ('serviço', 'título'):
                linhas = [d for d in grupo if d.tipo == tipo]
                so = sum((d.oficial or Decimal(0)) for d in linhas)
                se = sum((d.exportado for d in linhas), Decimal(0))
                p(f'    {tipo}: {len(linhas)} linha(s) com diferença; soma do AV oficial {so}, soma do saldo exportado '
                  f'{se}, diferença {so - se}')
            tot = [d for d in grupo if d.tipo == 'total']
            for d in tot:
                p(f'    total: oficial {d.oficial} x exportado {d.exportado} ({d.motivo})')
        for d in grupo:
            p(f'    oficial linha {d.linha_oficial} col {d.coluna_oficial} | export linha {d.linha_exportada} '
              f'col {d.coluna_exportada} | {d.codigo} | {d.campo}: oficial {_v(d.oficial)!s} x '
              f'exportado {_v(d.exportado)!s} ({d.motivo})')
    p('')
    p(f'Diferenças NÃO explicadas: {len(res.nao_explicadas)}')
    for d in res.nao_explicadas:
        p(f'  oficial linha {d.linha_oficial} col {d.coluna_oficial} | export linha {d.linha_exportada} '
          f'col {d.coluna_exportada} | {d.codigo} | {d.campo}: oficial {_v(d.oficial)} x '
          f'exportado {_v(d.exportado)} ({d.motivo})')
    p('')
    if res.nao_explicadas:
        p(f'RESULTADO: {len(res.nao_explicadas)} diferença(s) não explicada(s). Mostrar ao Tiago; '
          'não ajustar regra para fechar.')
    else:
        p(f'RESULTADO: {res.linhas_casadas} linhas casadas, 0 diferença não explicada.')
    return '\n'.join(out) + '\n'


def main(argv=None):
    ap = argparse.ArgumentParser(description='Confere o export do boletim contra a planilha oficial do Lote 09.')
    ap.add_argument('oficial')
    ap.add_argument('exportado')
    ap.add_argument('--relatorio', default=os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                                        '_retrato', 'conferencia_export.txt'))
    args = ap.parse_args(argv)
    res = conferir(args.oficial, args.exportado)
    texto = formatar_relatorio(res)
    sys.stdout.write(texto)
    os.makedirs(os.path.dirname(os.path.abspath(args.relatorio)), exist_ok=True)
    with open(args.relatorio, 'w', encoding='utf-8') as f:
        f.write(texto)
    print(f'relatório gravado em {args.relatorio}')
    return 1 if res.nao_explicadas else 0


if __name__ == '__main__':
    sys.exit(main())
