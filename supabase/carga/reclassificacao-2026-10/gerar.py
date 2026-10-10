"""Gera a planilha de-para da reclassificacao a partir de de-para.sql (so leitura).

Uso: python3 -I gerar.py <saida.xlsx>   (rodar de dentro do clone com supabase/.temp linkado)
"""
import json, subprocess, sys, pathlib
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

AQUI = pathlib.Path(__file__).resolve().parent
RAIZ = AQUI.parents[2]
saida = pathlib.Path(sys.argv[1])
bruto = subprocess.run(["npx", "supabase", "db", "query", "--linked", "-o", "json", "-f", str(AQUI / "de-para.sql")],
                       cwd=RAIZ, capture_output=True, text=True, check=True).stdout
d = json.loads(bruto[bruto.index("{"):]) if bruto.lstrip().startswith("{") or "{" in bruto else json.loads(bruto)
linhas = d["rows"] if isinstance(d, dict) and "rows" in d else d

REGRAS = {
    "a_centro_socio": "Centro do sócio (James PF, Casa James) → Distribuição a sócio",
    "a_despesa_pessoal": "Despesa pessoal da família / envio a PF → Distribuição a sócio",
    "b_empresa_ligada": "Amazônia / Juruá FM → Mútuo (enviado ou devolvido)",
    "c_consorcio": "Consórcio como despesa → Consórcio a contemplar",
    "c_equipamento": "Paccar / Hilux / roll-on em Outras despesas → Aquisição de Equipamento",
    "c_terreno": "Terreno em Outras despesas → Compra de Terreno",
    "d_imobilizado": "Centro de aquisição com categoria de custo → Aquisição de Equipamento",
    "d_imobilizado_gasto": "Frete/documentação no centro de aquisição → DECIDIR (Manutenção/Doc?)",
    "e_pro_labore_sem_socio": "Pró-labore sem nome do sócio → DECIDIR de quem é",
    "e_pecuaria": "Custeio de pecuária → DECIDIR",
    "e_despesas_financeiras": "Despesas financeiras genéricas → DECIDIR",
}
COLS = ["rateio_id", "numero", "tipo", "competencia", "valor", "descricao", "fornecedor", "categoria_atual",
        "centro_atual", "centro_raiz", "categoria_proposta", "centro_proposto", "regra", "decidir"]
TITULOS = ["Rateio (id)", "Lançamento", "Tipo", "Competência", "Valor", "Descrição", "Fornecedor", "Categoria atual",
           "Centro atual", "Centro raiz", "Categoria proposta", "Centro proposto", "Regra", "Decidir?", "Aprovar (sim/não/outra)"]

wb = Workbook()
res = wb.active; res.title = "Resumo"
res.append(["Regra", "O que faz", "Rateios", "Valor (R$)", "Para decidir"]); 
por = defaultdict(lambda: [0, 0.0, 0])
for l in linhas:
    p = por[l["regra"]]; p[0] += 1; p[1] += float(l["valor"]); p[2] += l["decidir"] == "sim"
for regra in REGRAS:
    if regra in por:
        q, v, dec = por[regra]; res.append([regra, REGRAS[regra], q, round(v, 2), dec])
res.append(["TOTAL", "", sum(p[0] for p in por.values()), round(sum(p[1] for p in por.values()), 2), sum(p[2] for p in por.values())])

def aba(nome, filtro):
    ws = wb.create_sheet(nome); ws.append(TITULOS)
    for l in linhas:
        if filtro(l):
            ws.append([float(l[c]) if c == "valor" else l[c] for c in COLS] + [""])
    for i, _ in enumerate(TITULOS, 1):
        ws.column_dimensions[get_column_letter(i)].width = 18
    ws.column_dimensions["F"].width = 50
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for c in ws[1]: c.font = Font(bold=True)

aba("De-para", lambda l: True)
aba("Decidir", lambda l: l["decidir"] == "sim")
for c in res[1]: c.font = Font(bold=True)
res.column_dimensions["B"].width = 70
saida.parent.mkdir(parents=True, exist_ok=True)
wb.save(saida)
print(f"{len(linhas)} rateios, {saida}")
for regra, (q, v, dec) in sorted(por.items()):
    print(f"  {regra}: {q} rateios, R$ {v:,.2f}, decidir {dec}")
