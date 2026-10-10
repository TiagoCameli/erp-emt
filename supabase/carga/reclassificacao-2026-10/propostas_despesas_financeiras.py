"""Proposta de categoria para cada rateio em "Despesas financeiras" (Tiago, 10/10/2026).

Le os rateios da regra e_despesas_financeiras de de-para.sql e propoe a categoria
por palavra-chave da descricao. Confianca "alta" = palavra inequivoca; "baixa" =
pista fraca; "sem pista" = so o Tiago sabe. So leitura.
Uso: python3 propostas_despesas_financeiras.py <saida.xlsx>
"""
import json, re, subprocess, sys, pathlib
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

AQUI = pathlib.Path(__file__).resolve().parent
RAIZ = AQUI.parents[2]
bruto = subprocess.run(["npx", "supabase", "db", "query", "--linked", "-o", "json", "-f", str(AQUI / "de-para.sql")],
                       cwd=RAIZ, capture_output=True, text=True, check=True).stdout
d = json.loads(bruto[bruto.index("{"):])
linhas = [l for l in (d["rows"] if isinstance(d, dict) else d) if l["regra"] == "e_despesas_financeiras"]

# (padrao, categoria, confianca, observacao). A primeira que casar vale.
REGRAS = [
    (r"TRANSFER[EÊ]NCIA ENTRE CONTAS", None, "alta", "Não é despesa: é transferência entre contas lançada como pagamento. Excluir o lançamento e registrar como transferência"),
    (r"HILUX", "Aquisição de Equipamento", "alta", "Compra do veículo"),
    (r"SEGURO|AP[OÓ]LICE|SEGURADORA", "Seguros", "alta", ""),
    (r"IPVA", "IPVA", "alta", ""),
    (r"DETRAN|LICENCIAMENTO|ANTT|ANUIDADE|CERTIFICADO", "Impostos e taxas", "alta", ""),
    (r"CART[OÓ]RIO", "Cartório", "alta", ""),
    (r"ADVOGAD|JUR[IÍ]DIC|PARECER|RECURSO LICITA", "Jurídico", "alta", ""),
    (r"IRPF|RECEITA ?FEDERAL", "Impostos", "baixa", "IRPF é imposto de pessoa física: se for do sócio, é Distribuição a sócio"),
    (r"FARDA|FARDAMENTO|BLUSAS", "EPI'S", "alta", ""),
    (r"EXAME|ADMISSIONA|DEMISSIONA|OCUPACIONA|SEGURAN[CÇ]A DO TRABALHO|\bSST\b|\bPGR\b|SESI", "Escritório e administrativo", "alta", "Saúde e segurança do trabalho"),
    (r"PLANO .*SA[UÚ]DE|BRADESCO SAUDE|UNIMED", "Folha de pagamento", "baixa", "Benefício de funcionário (se for de família do sócio, é Distribuição)"),
    (r"INTERNET", "Internet", "alta", ""),
    (r"NASAJON|SOFTWARE|DESENVOLVIMENTO SOFT|APLICATIVO|SISTEMA DE GEST", "Escritório e administrativo", "alta", "Sistema/software"),
    (r"MONITORAMENTO|SISTEMA DE SEGURAN|CONCERTINA", "Escritório e administrativo", "alta", "Segurança patrimonial"),
    (r"NOTEBOOK|INFORM[AÁ]TICA|CELULAR|SMARTPHONE", "Escritório e administrativo", "alta", "Equipamento de escritório"),
    (r"AR[ -]?COND|COLCH|\bTV\b|CONTAINER", "Materiais", "baixa", "Bem para alojamento da obra"),
    (r"FEIRA|ALIMENTA|COZINHA|AUX[IÍ]LIO ALIMENTA", "Vale Alimentação Mão de Obra", "alta", ""),
    (r"ALUGUEL", "Aluguel", "alta", ""),
    (r"FRETE", "Frete", "alta", ""),
    (r"ABASTECIMENTO|COMBUST", "Combustível", "alta", ""),
    (r"JUROS", "Juros", "alta", ""),
    (r"LICITA", "Escritório e administrativo", "baixa", "Despesa de licitação"),
    (r"RETIRA.*SR\.? ?JAMES|DESPESAS DO NETO", "Distribuição a sócio", "baixa", "Dinheiro do sócio/família"),
    (r"R[AÁ]DIO E TV JURU", "Mútuo concedido a empresa ligada", "baixa", "Juruá FM é empresa ligada"),
]

def propor(desc):
    t = (desc or "").upper()
    for padrao, cat, conf, obs in REGRAS:
        if re.search(padrao, t):
            return cat, conf, obs
    return None, "sem pista", "Descrição genérica (despesas, acerto, pagamento, importado do Mais Controle): só você sabe"

wb = Workbook(); ws = wb.active; ws.title = "Propostas"
cab = ["Rateio (id)", "Lançamento", "Competência", "Valor", "Descrição", "Fornecedor", "Centro atual",
       "Categoria proposta", "Confiança", "Observação", "Aprovar (sim/não/outra)"]
ws.append(cab)
resumo = defaultdict(lambda: [0, 0.0])
for l in linhas:
    cat, conf, obs = propor(l["descricao"])
    ws.append([l["rateio_id"], l["numero"], l["competencia"], float(l["valor"]), l["descricao"], l["fornecedor"],
               l["centro_atual"], cat or "", conf, obs, ""])
    k = (cat or "(sem proposta)", conf); resumo[k][0] += 1; resumo[k][1] += float(l["valor"])
for i in range(1, len(cab) + 1): ws.column_dimensions[get_column_letter(i)].width = 18
ws.column_dimensions["E"].width = 55; ws.column_dimensions["J"].width = 55
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
for c in ws[1]: c.font = Font(bold=True)
rs = wb.create_sheet("Resumo", 0); rs.append(["Categoria proposta", "Confiança", "Rateios", "Valor (R$)"])
for (cat, conf), (q, v) in sorted(resumo.items(), key=lambda x: -x[1][1]): rs.append([cat, conf, q, round(v, 2)])
for c in rs[1]: c.font = Font(bold=True)
rs.column_dimensions["A"].width = 38
saida = pathlib.Path(sys.argv[1]); wb.save(saida)
print(len(linhas), "rateios ->", saida)
for (cat, conf), (q, v) in sorted(resumo.items(), key=lambda x: -x[1][1]): print(f"  {cat} [{conf}]: {q}, R$ {v:,.2f}")
