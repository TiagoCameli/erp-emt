"""Planilha do que ficou por fazer na reclassificacao (so leitura).

Roda de-para.sql de novo (o que ja foi aplicado nao casa mais com nenhuma
regra), tira as linhas que o Tiago marcou "nao" na planilha anterior (decididas:
ficam como estao), e junta a proposta de categoria das "Despesas financeiras".
Uso: python3 gerar_pendentes.py <de-para do Tiago.xlsx> <saida.xlsx> [pendentes marcadas.xlsx ...]
(as planilhas de pendentes ja marcadas entram pelo id do rateio: "nao" fica de fora)
"""
import json, subprocess, sys, pathlib
from collections import defaultdict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

AQUI = pathlib.Path(__file__).resolve().parent
RAIZ = AQUI.parents[2]

def rodar(sql):
    out = subprocess.run(["npx", "supabase", "db", "query", "--linked", "-o", "json", "-f", str(sql)],
                         cwd=RAIZ, capture_output=True, text=True, check=True).stdout
    d = json.loads(out[out.index("{"):]); return d["rows"] if isinstance(d, dict) else d

# A funcao propor() do script de propostas, sem rodar o script inteiro.
fonte = (AQUI / "propostas_despesas_financeiras.py").read_text(encoding="utf-8")
ns = {"re": __import__("re")}; exec(fonte[fonte.index("# (padrao, categoria"):fonte.index("wb = Workbook()")], ns)
propor = ns["propor"]

linhas = rodar(AQUI / "de-para.sql")

# Linhas "nao" da planilha do Tiago: casar por lancamento, valor, centro e descricao.
wb_t = load_workbook(sys.argv[1], data_only=True); ws_t = wb_t["De-para"]
hdr = [c.value for c in ws_t[1]]
ap = next(h for h in hdr if h and h.startswith("Aprovar"))
chave = lambda n, v, c, d: (n, round(float(v), 2), c, (d or "").strip())
nao = defaultdict(int)
for r in ws_t.iter_rows(min_row=2, values_only=True):
    r = dict(zip(hdr, r))
    if str(r[ap] or "").strip().lower() == "não":
        nao[chave(r["Lançamento"], r["Valor"], r["Centro atual"], r["Descrição"])] += 1

OBS = {
    "a_centro_socio": "Frete de peças: no Mais Controle a parte estava em 'Casa James' sem equipamento. Diga o equipamento ou deixe como distribuição",
    "a_despesa_pessoal": "Achado pelo texto (envio a PF, despesa pessoal). Atenção: a regra pegou nomes como DONIZETE por engano",
    "b_empresa_ligada": "Empréstimo (BASA) em centro da Amazônia: sugiro deixar como está",
    "c_consorcio": "Consórcio lançado como despesa",
    "c_equipamento": "Bem comprado lançado em Outras despesas",
    "d_imobilizado": "Centro de aquisição com categoria de custo",
    "d_imobilizado_gasto": "Frete/documentação no centro de aquisição: é gasto (Manutenção/Doc.) ou parte do bem?",
    "e_pecuaria": "Custeio de pecuária",
    "e_pro_labore_sem_socio": "Pró-labore sem nome do sócio",
}
# Linhas ja decididas nas planilhas de pendentes marcadas, pelo id do rateio.
nao_ids = set()
for extra in sys.argv[3:]:
    wb_e = load_workbook(extra, data_only=True)
    for aba in wb_e.sheetnames:
        ws_e = wb_e[aba]; h = [c.value for c in ws_e[1]]
        if not h or not str(h[0] or "").startswith("Rateio"): continue
        col = next((i for i, x in enumerate(h) if x and str(x).startswith("Aprovar")), None)
        col_o = next((i for i, x in enumerate(h) if x and str(x).startswith("Se outra")), None)
        for r in ws_e.iter_rows(min_row=2, values_only=True):
            # Decidida (sim, nao ou outra) sai da planilha: sim/outra ja foram aplicadas, nao fica como esta.
            marcada = col is not None and str(r[col] or "").strip() != ""
            outra = col_o is not None and str(r[col_o] or "").strip() != ""
            if marcada or outra: nao_ids.add(r[0])

pend = []
for l in linhas:
    k = chave(l["numero"], l["valor"], l["centro_atual"], l["descricao"])
    if nao.get(k):
        nao[k] -= 1; continue
    if l["rateio_id"] in nao_ids:
        continue
    cat, conf, obs = l["categoria_proposta"], "", OBS.get(l["regra"], "")
    if l["regra"] == "e_despesas_financeiras":
        cat, conf, obs = propor(l["descricao"])
    pend.append({**l, "categoria_proposta": cat or "", "confianca": conf, "obs": obs})

wb = Workbook(); rs = wb.active; rs.title = "Resumo"
rs.append(["Regra", "Pendentes", "Valor (R$)", "O que falta"])
por = defaultdict(lambda: [0, 0.0])
for p in pend: por[p["regra"]][0] += 1; por[p["regra"]][1] += float(p["valor"])
for regra, (q, v) in sorted(por.items()): rs.append([regra, q, round(v, 2), OBS.get(regra, "Despesas financeiras: proposta por palavra-chave, confira a coluna Confiança")])
rs.append(["TOTAL", len(pend), round(sum(float(p["valor"]) for p in pend), 2), ""])
rs.append([]); rs.append(["Já aplicado em 10/10/2026: parte 1 (443 'sim'), parte 2 (13 'outra'), parte 3 (Amazônia 463, Colorado, Areacre, apólice), parte 4 (61 da planilha de pendentes). Os 'não' ficam como estão e saem da planilha."])
ws = wb.create_sheet("Pendentes")
cab = ["Rateio (id) — não apagar", "Lançamento", "Tipo", "Competência", "Valor", "Descrição", "Fornecedor", "Categoria atual",
       "Centro atual", "Categoria proposta", "Centro proposto", "Regra", "Confiança", "Observação", "Aprovar (sim/não/outra)", "Se outra: categoria / centro certos"]
ws.append(cab)
for p in pend:
    ws.append([p["rateio_id"], p["numero"], p["tipo"], p["competencia"], float(p["valor"]), p["descricao"], p["fornecedor"],
               p["categoria_atual"], p["centro_atual"], p["categoria_proposta"], p["centro_proposto"], p["regra"],
               p["confianca"], p["obs"], "", ""])
for i in range(1, len(cab) + 1): ws.column_dimensions[get_column_letter(i)].width = 18
ws.column_dimensions["F"].width = 55; ws.column_dimensions["N"].width = 50
ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions
for c in ws[1]: c.font = Font(bold=True)
for c in rs[1]: c.font = Font(bold=True)
rs.column_dimensions["D"].width = 90
# Pagos a terceiros citando o Tiago, que as regras do de-para nao pegam.
import re
ids_pend = {p["rateio_id"] for p in pend}
tiago = [t for t in rodar(AQUI / "em-nome-do-tiago.sql")
         if t["grupo"] == "terceiro" and t["natureza"] not in ("distribuicao", "mutuo")
         and t["rateio_id"] not in ids_pend and t["rateio_id"] not in nao_ids]
wt = wb.create_sheet("Em nome do Tiago")
cab_t = ["Rateio (id) — não apagar", "Lançamento", "Competência", "Valor", "Descrição", "Fornecedor", "Categoria atual",
         "Centro atual", "Sugestão", "Observação", "Aprovar (sim/não/outra)", "Se outra: categoria / centro certos"]
wt.append(cab_t)
for t in tiago:
    d = (t["descricao"] or "").upper()
    if "RETIRADA" in d: sug, obs = "Distribuição a sócio / Sócio Tiago de Melo Cameli", "Retirada sua: é distribuição"
    elif re.search(r"ABASTEC|LAVAGEM|HILUX|MOTORISTA", d): sug, obs = "Deixar como está", "Carro/motorista em uso na obra: custo da obra, salvo uso pessoal"
    elif "PASSAGENS" in d: sug, obs = "Decidir", "Viagem a trabalho (fica) ou pessoal (distribuição)?"
    elif "ESCRITURA" in d: sug, obs = "Decidir", "Escritura e inventário da fazenda, pago a James, hoje como empréstimo em Aquisição de Imóveis"
    else: sug, obs = "Decidir", "A observação do lançamento cita Tiago; descrição genérica"
    wt.append([t["rateio_id"], t["numero"], t["comp"], float(t["valor"]), t["descricao"], t["fornecedor"], t["categoria"], t["centro"], sug, obs, "", ""])
for i in range(1, len(cab_t) + 1): wt.column_dimensions[get_column_letter(i)].width = 18
wt.column_dimensions["E"].width = 55; wt.column_dimensions["J"].width = 60
wt.freeze_panes = "B2"; wt.auto_filter.ref = wt.dimensions
for c in wt[1]: c.font = Font(bold=True)
rs.append(["Em nome do Tiago (aba própria)", len(tiago), round(sum(float(t["valor"]) for t in tiago), 2), "Pagos a terceiros citando Tiago (combustível da Hilux, passagens, retiradas)"])
wb.save(sys.argv[2])
print(len(pend), "pendentes ->", sys.argv[2])
for regra, (q, v) in sorted(por.items()): print(f"  {regra}: {q}, R$ {v:,.2f}")
print("nao nao casados:", sum(v for v in nao.values() if v > 0))
